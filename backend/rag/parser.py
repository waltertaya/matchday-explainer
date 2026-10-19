from docx import Document as DocxDocument
from pptx import Presentation
import pdfplumber
import os

from typing import Dict, Tuple

class Parser:
    def __init__(self, path: str):
        self.path = path

    def parse_txt(self) -> str:
        ''' extract text from txt files '''
        with open(self.path, 'r', encoding='utf-8') as file:
            return file.read()

    def parse_pdf(self) -> str:
        ''' extract text from pdf file '''
        text = []

        with pdfplumber.open(self.path) as pdf:
            for page in pdf.pages:
                t = page.extract_text()
                if t:
                    text.append(t)

        return '\n\n'.join(text)

    def parse_docx(self) -> str:
        ''' extract text from word document '''
        doc = DocxDocument(self.path)
        paragraphs = []

        for paragraph in doc.paragraphs:
            if paragraph.text and paragraph.text.strip():
                paragraphs.append(paragraph.text)

        return '\n\n'.join(paragraphs)

    def parse_pptx(self) -> str:
        ''' extract text from pptx files '''
        prs = Presentation(self.path)
        text = []

        for slide in prs.slides:
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for para in shape.text_frame.paragraphs:
                        for run in para.runs:
                            text.append(run.text)

        return '\n\n'.join(text)


    def parse(self) -> Tuple[str, Dict]:
        ''' extract text from given document '''
        ext = os.path.splitext(self.path)[1].lower()

        if ext in ['.txt', '.csv', '.md']:
            text = self.parse_txt()
        elif ext == '.pdf':
            text = self.parse_pdf()
        elif ext == '.docx':
            text = self.parse_docx()
        elif ext == '.pptx':
            text = self.parse_pptx()
        else:
            raise ValueError(f'Unsupported file type: {ext}. Supported files - .txt, .csv, .md, .pdf, .docx, .pptx')

        metadata = {
            'source': os.path.basename(self.path),
            'file_type': ext.lstrip('.'),
            'path': self.path,
            'size_bytes': os.path.getsize(self.path)
        }

        return text, metadata
